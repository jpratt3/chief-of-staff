"""Acceptance tests. Mirrors the criteria in the build prompts.

The browser tests drive the real mock portals through Playwright, so they are
slow but they are the ones that actually prove the agent works.
"""
from __future__ import annotations

import asyncio
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config          # noqa: E402
import extract         # noqa: E402
import orchestrator as orch  # noqa: E402
import reconcile as recon    # noqa: E402
import seed            # noqa: E402
import services        # noqa: E402
import summary as summary_mod  # noqa: E402
from agent import chase as chase_mod  # noqa: E402
from agent import executor, mailbox   # noqa: E402

TARGET = {"claims": 47, "paid": 412_655.00, "incurred": 538_190.00}


@pytest.fixture(scope="session", autouse=True)
def world():
    services.start()
    seed.generate_all()
    yield


@pytest.fixture(scope="session")
def account():
    import httpx
    a = httpx.get(f"{config.epic_url()}/accounts/meridian", timeout=10).json()
    a["expires"] = "2026-03-01"
    return a


def _policy(account, carrier):
    return next(p for p in account["policies"] if p["carrier"] == carrier)


# ---- seed -----------------------------------------------------------------

def test_seed_hits_hero_card_totals():
    t = seed.book_totals()
    assert t["claims"] == TARGET["claims"]
    assert t["paid"] == pytest.approx(TARGET["paid"], abs=0.01)
    assert t["incurred"] == pytest.approx(TARGET["incurred"], abs=0.01)
    assert t["policies"] == 12


def test_all_six_layouts_generate():
    made = seed.generate_all()
    assert len(made) == 8  # 6 carriers + large-loss detail + Summit prior


# ---- extraction -----------------------------------------------------------

def test_every_layout_extracts_to_schema():
    ex = extract.LayoutExtractor()
    for p in config.SEED_PDF_DIR.glob("*.pdf"):
        doc = ex.extract(p)
        assert doc.carrier, f"{p.name}: no carrier parsed"
        if p.name.startswith("BLH"):
            continue  # Bluehaven policy is never retrieved, so it carries no claims
        assert doc.claims, f"{p.name}: no claims parsed"
        for c in doc.claims:
            assert c.source_ref.artifact == p.name
            assert c.source_ref.page >= 1


def test_current_loss_runs_reconcile_to_the_hero_card():
    ex = extract.LayoutExtractor()
    names = ["NBM_LossRun_2021-2025.pdf", "HBS_LossRun_2021-2025.pdf",
             "RDS_LossRun_2021-2025.pdf", "SMT_LossRun_2021-2025.pdf",
             "IRN_LossRun_2021-2025.pdf"]
    docs = [ex.extract(config.SEED_PDF_DIR / n) for n in names]
    assert sum(len(d.claims) for d in docs) == TARGET["claims"]
    assert sum(d.total_paid for d in docs) == pytest.approx(TARGET["paid"], abs=0.01)
    assert sum(d.total_incurred for d in docs) == pytest.approx(TARGET["incurred"], abs=0.01)


# ---- reconciliation -------------------------------------------------------

@pytest.fixture(scope="module")
def reconciliation():
    ex = extract.LayoutExtractor()
    names = ["NBM_LossRun_2021-2025.pdf", "HBS_LossRun_2021-2025.pdf",
             "HBS_LargeLoss_Detail.pdf", "RDS_LossRun_2021-2025.pdf",
             "SMT_LossRun_2021-2025.pdf", "SMT_LossRun_2025_Prior.pdf",
             "IRN_LossRun_2021-2025.pdf"]
    docs = [ex.extract(config.SEED_PDF_DIR / n) for n in names]
    return docs, recon.reconcile(docs)


def test_exactly_three_discrepancies_and_no_false_positives(reconciliation):
    _docs, rec = reconciliation
    assert len(rec.discrepancies) == 3, [d.claim_number for d in rec.discrepancies]
    kinds = {d.kind for d in rec.discrepancies}
    assert kinds == {"amount_conflict", "date_conflict", "missing_claim"}


def test_reserve_development_is_annotated_not_flagged(reconciliation):
    _docs, rec = reconciliation
    assert len(rec.development) == 1
    note = rec.development[0]
    assert note.from_value == pytest.approx(18_200.0)
    assert note.to_value == pytest.approx(41_700.0)
    assert note.claim_number not in {d.claim_number for d in rec.discrepancies}


def test_amount_conflict_matches_the_reference_case(reconciliation):
    _docs, rec = reconciliation
    d = next(x for x in rec.discrepancies if x.kind == "amount_conflict")
    assert d.claim_number == "GL-0088314"
    assert d.pick_value == "$164,000"
    assert d.other_value == "$158,500"
    assert d.pick_ref.artifact == "HBS_LossRun_2021-2025.pdf"
    assert d.other_ref.artifact == "HBS_LargeLoss_Detail.pdf"
    assert d.pick_label == "Carrier loss run"


def test_canonical_set_matches_the_hero_card(reconciliation):
    _docs, rec = reconciliation
    t = rec.totals
    assert t["claims"] == TARGET["claims"]
    assert t["paid"] == pytest.approx(TARGET["paid"], abs=0.01)
    assert t["incurred"] == pytest.approx(TARGET["incurred"], abs=0.01)


def test_every_discrepancy_cites_a_resolvable_page(reconciliation):
    _docs, rec = reconciliation
    for d in rec.discrepancies:
        for ref in (d.pick_ref, d.other_ref):
            assert (config.SEED_PDF_DIR / ref.artifact).exists()
            assert ref.page >= 1


# ---- deliverable ----------------------------------------------------------

def test_workbook_builds_with_correct_totals(reconciliation, tmp_path):
    from openpyxl import load_workbook
    docs, rec = reconciliation
    out = summary_mod.build_workbook(
        "Meridian Logistics", rec, docs, tmp_path / "s.xlsx",
        seed.VALUATION_CURRENT, (seed.PERIOD_START, seed.PERIOD_END), 12)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "By Year", "Large Losses", "Sources"]
    ws = wb["Summary"]
    last = ws.max_row
    assert ws.cell(row=last, column=6).value == pytest.approx(TARGET["paid"])
    assert ws.cell(row=last, column=7).value == pytest.approx(TARGET["incurred"])


# ---- chase ----------------------------------------------------------------

def test_chase_cadence_and_wrong_policy_rejection():
    mailbox.reset()
    ch = chase_mod.Chase(policy_number="IW-PR-440255", carrier="Ironwood",
                         account="Meridian Logistics")
    for day in range(1, 15):
        ch.tick(day)
    assert ch.rejected, "the wrong-policy attachment should have been rejected"
    assert ch.state == chase_mod.RECEIVED
    assert any("follow-up 1" in l for l in ch.log)
    assert any("follow-up 2" in l for l in ch.log)
    sent = mailbox.sent_for("IW-PR-440255")
    days = {m["kind"]: m["virtual_day"] for m in sent}
    assert days.get("followup1") == config.FOLLOWUP_1_DAY
    assert days.get("followup2") == config.FOLLOWUP_2_DAY


def test_summit_returns_two_valuations():
    mailbox.reset()
    ch = chase_mod.Chase(policy_number="SM-AU-330190", carrier="Summit",
                         account="Meridian Logistics")
    for day in range(1, 15):
        ch.tick(day)
    assert ch.state == chase_mod.RECEIVED
    assert len(ch.attachments) == 2


# ---- agent (browser) ------------------------------------------------------

def test_redstone_self_corrects_the_policy_format(account):
    pol = _policy(account, "Redstone")
    res = asyncio.run(executor.run_policy(account, pol))
    assert res.status == "retrieved"
    events = [json.loads(l) for l in
              (res.run_dir / "trace.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    fills = [e for e in events if e["tool"] == "browser_fill"
             and e["args"].get("selector") == "#policy"]
    tried = [f["args"]["value"] for f in fills]
    assert "RS-221-0567" in tried, "should try the AMS-formatted number first"
    assert "RS2210567" in tried, "should retry reformatted after reading the page"
    rejected = [e for e in events if "Invalid policy number format" in e["result"]]
    assert rejected, "the rejection must be observed, not assumed"


def test_bluehaven_escalates_without_guessing(account):
    pol = _policy(account, "Bluehaven")
    res = asyncio.run(executor.run_policy(account, pol))
    assert res.status == "input_required"
    assert not res.artifacts
    events = [json.loads(l) for l in
              (res.run_dir / "trace.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    answers = [e for e in events if e["tool"] == "browser_fill"
               and e["args"].get("selector") == "#answer"]
    assert not answers, "the agent must not attempt the security question"


def test_harborstone_takes_both_documents(account):
    pol = _policy(account, "Harborstone")
    res = asyncio.run(executor.run_policy(account, pol))
    assert res.status == "retrieved"
    assert sorted(res.artifacts) == ["HBS_LargeLoss_Detail.pdf", "HBS_LossRun_2021-2025.pdf"]


def test_northbridge_clears_totp(account):
    pol = _policy(account, "Northbridge")
    res = asyncio.run(executor.run_policy(account, pol))
    assert res.status == "retrieved"
    events = [json.loads(l) for l in
              (res.run_dir / "trace.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    assert any(e["args"].get("secret_ref", "").endswith("totp_code")
               for e in events if e["tool"] == "browser_fill")


def test_no_secret_ever_reaches_a_trace(account):
    """The acceptance criterion from the build prompt: grep proves it."""
    vault = json.loads(config.VAULT_PATH.read_text())
    secrets = {v for entry in vault.values() for k, v in entry.items()
               if k in ("password", "totp_secret") and v}
    assert secrets
    leaked = []
    for tp in config.RUNS_DIR.glob("*/trace.jsonl"):
        blob = tp.read_text(encoding="utf-8")
        for s in secrets:
            if s in blob:
                leaked.append((tp.name, s[:4] + "..."))
    assert not leaked, f"secret leaked into trace: {leaked}"
