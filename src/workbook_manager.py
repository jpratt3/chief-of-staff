from __future__ import annotations

from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook, load_workbook


WORKBOOK_FILENAME = "chief_of_staff_tracker.xlsx"
WORKBOOK_DIRNAME = "workbook"


SHEET_COLUMNS: Dict[str, List[str]] = {
    "Overview": [
        "section",
        "metric",
        "value",
        "window",
        "updated_at",
    ],
    "Client_Project_Summary": [
        "client_name",
        "project_name",
        "project_type",
        "renewal_date",
        "current_stage",
        "next_milestone",
        "next_milestone_date",
        "overall_status",
        "priority",
        "days_to_due",
        "last_activity_date",
        "open_work_item_count",
        "overdue_work_item_count",
        "blocked_flag",
        "owner",
        "notes",
    ],
    "Client_Project_Work_Items": [
        "client_name",
        "project_name",
        "work_item_name",
        "work_item_type",
        "status",
        "owner",
        "due_date",
        "completed_flag",
        "completed_at",
        "completion_source",
        "manual_override",
        "notes",
        "evidence_type",
        "evidence_reference",
        "last_activity_date",
        "blocked_flag",
        "stage_name",
        "priority",
        "project_type",
        "sort_order",
    ],
    "Workflow_Template": [
        "stage_name",
        "stage_window",
        "task_name",
        "role_owner_type",
        "ae_flag",
        "aae_flag",
        "ar_flag",
        "timing_detail",
        "stage_sort_order",
        "task_sort_order",
        "active",
        "template_source",
    ],
    "Activity_Log": [
        "activity_timestamp",
        "activity_date",
        "client_name",
        "project_name",
        "work_item_name",
        "activity_type",
        "source_type",
        "source_reference",
        "sender_or_organizer",
        "subject_or_title",
        "matched_by_rule",
        "notes",
    ],
    "Evidence_Rules": [
        "rule_name",
        "active",
        "project_type",
        "stage_name",
        "evidence_type",
        "match_field",
        "match_text",
        "suggested_status",
        "suggested_completed_flag",
        "confidence",
        "needs_manual_review",
        "notes",
    ],
    "Lookup_Lists": [
        "list_name",
        "value",
        "sort_order",
        "active",
    ],
    "Unmatched_Events": [
        "source_type",
        "timestamp",
        "subject",
        "folder_name",
        "discard_reason",
    ],
}


def get_workbook_path(project_root: Path) -> Path:
    workbook_dir = project_root / WORKBOOK_DIRNAME
    workbook_dir.mkdir(parents=True, exist_ok=True)
    return workbook_dir / WORKBOOK_FILENAME


def _create_new_workbook() -> Workbook:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for sheet_name, columns in SHEET_COLUMNS.items():
        worksheet = workbook.create_sheet(title=sheet_name)
        worksheet.append(columns)

    return workbook


def _ensure_sheet_headers(worksheet, columns: List[str]) -> None:
    current_headers = [cell.value for cell in worksheet[1]]

    if current_headers == columns:
        return

    if all(value is None for value in current_headers):
        for index, column_name in enumerate(columns, start=1):
            worksheet.cell(row=1, column=index, value=column_name)
        return

    worksheet.delete_rows(1, worksheet.max_row)
    worksheet.append(columns)


def ensure_workbook(project_root: Path) -> Path:
    workbook_path = get_workbook_path(project_root)

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = _create_new_workbook()

    for sheet_name, columns in SHEET_COLUMNS.items():
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(title=sheet_name)
            worksheet.append(columns)
        else:
            worksheet = workbook[sheet_name]
            _ensure_sheet_headers(worksheet, columns)

    for existing_sheet in list(workbook.sheetnames):
        if existing_sheet not in SHEET_COLUMNS:
            del workbook[existing_sheet]

    workbook.save(workbook_path)
    workbook.close()
    return workbook_path
